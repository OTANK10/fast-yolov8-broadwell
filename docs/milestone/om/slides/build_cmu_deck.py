#!/usr/bin/env python3
"""Om's mid-project talk on the CMU Digital Tartan template (cmu_template.pptx).

    python docs/milestone/om/slides/build_cmu_deck.py   ->  om_mid_presentation_cmu.pptx

Slide 1 uses the template's "Title Slide" layout (red tartan); every other slide uses
"Title Only", so the template's title style, tartan band and wordmark come from its
layouts. Content stays above the band (y <= 4.6 in). Same content and numbers as
build_mid_deck.js; every % of peak assumes 3.4 GHz (Assignment 2 convention).
"""
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "../../../.."))
TEMPLATE = os.path.join(HERE, "cmu_template.pptx")
OUT = os.path.join(HERE, "om_mid_presentation_cmu.pptx")
fig = lambda p: os.path.join(ROOT, p)

RED, INK, MUTED, TEAL, TEALSOFT = RGBColor(0xBB, 0x00, 0x27), RGBColor(0x1B, 0x26, 0x30), RGBColor(0x5A, 0x68, 0x75), RGBColor(0x0C, 0x6A, 0x87), RGBColor(0xD7, 0xEA, 0xF0)
SOFT, GREY, AMBER, WHITE, PINK = RGBColor(0xF1, 0xF3, 0xF5), RGBColor(0x6B, 0x77, 0x82), RGBColor(0xB7, 0x73, 0x1A), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xF8, 0xE5, 0xE9)
AMBER_SOFT = RGBColor(0xF5, 0xE6, 0xCF)
FONT = "Open Sans"  # the template's body font

prs = Presentation(TEMPLATE)
# drop the template's two sample slides
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)
L_TITLE, L_ONLY = prs.slide_layouts[0], prs.slide_layouts[2]


def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, name=None):
    """runs: a string, or a list of paragraphs, each a string or a list of (text, {opts})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
        if i: p.space_before = Pt(o.get("space", 6) if not isinstance(para, str) else 6)
    return tb


def card(slide, x, y, w, h, fill=SOFT, name=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.08
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if name: s.name = name
    return s


def content_slide(title, notes):
    s = prs.slides.add_slide(L_ONLY)
    s.shapes.title.text = title
    s.notes_slide.notes_text_frame.text = notes
    return s


def stat(slide, x, y, w, big, small, color=TEAL):
    text(slide, x, y, w, 0.6, big, size=30, bold=True, color=color)
    text(slide, x, y + 0.62, w, 0.6, small, size=12, color=MUTED)


def style_chart(chart, fmt, color):
    chart.has_legend = False
    plot = chart.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.font.size, dl.font.name, dl.font.color.rgb = Pt(11), FONT, INK
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.series[0].format.fill.solid(); plot.series[0].format.fill.fore_color.rgb = color
    for ax in (chart.category_axis, chart.value_axis):
        ax.tick_labels.font.size, ax.tick_labels.font.name, ax.tick_labels.font.color.rgb = Pt(10), FONT, MUTED
        ax.format.line.fill.background()
    chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xDD, 0xE3, 0xE7)
    chart.has_title = True
    chart.chart_title.text_frame.paragraphs[0].runs or chart.chart_title.text_frame.paragraphs[0].add_run()


def chart_title(chart, t):
    tf = chart.chart_title.text_frame
    tf.text = t
    f = tf.paragraphs[0].runs[0].font
    f.size, f.name, f.bold, f.color.rgb = Pt(12), FONT, False, INK


# ------------------------------------------------------------------ 1. title
s = prs.slides.add_slide(L_TITLE)
text(s, 2.33, 1.5, 6.9, 1.3, [[("The 3×3 stride-1 convolution kernel", {})]], size=34, bold=True, color=WHITE, anchor=MSO_ANCHOR.BOTTOM, name="title")
text(s, 2.33, 3.0, 6.9, 0.5, [[("Om Tank", {"bold": True}), ("  ·  YOLOv8n on one Broadwell core, built like the Assignment 2 kernel", {})]], size=15, color=WHITE)
for i, (big, small) in enumerate([("32", "layers"), ("66%", "of YOLOv8n's work"), ("85%", "of peak, kernel alone"), ("77%", "of peak, 32 layers")]):
    text(s, 2.33 + i * 1.75, 3.85, 1.65, 0.55, big, size=26, bold=True, color=WHITE)
    text(s, 2.33 + i * 1.75, 4.4, 1.65, 0.4, small, size=11, color=WHITE)
s.notes_slide.notes_text_frame.text = (
    "My kernel is the 3x3 stride-1 convolution: 32 layers, the C2f bottlenecks and the Detect head, two thirds of YOLOv8n's work. "
    "The story: it is the Assignment 2 kernel, I followed the same four steps, and the measurements chose the tile. "
    "The 2x5 microkernel alone: 85% of the 32 FLOPs/cycle peak; inside all 32 real layers: 77%, 1.24 times faster than OpenCV. Intrinsics only, no assembly. All percentages assume the core runs at 3.4 GHz, as in Assignment 2.")

# ------------------------------------------------------------------ 2. same as A2
s = content_slide("It is the Assignment 2 kernel, with floats",
    "Same machine as Assignment 2 and the same kernel. In A2, m was 5 rows of A broadcast, n was 8 doubles = 2 vectors of B. "
    "Here m is 5 output pixels, n is 16 output channels = 2 vectors of 8 floats, and k is input channels times the 9 taps of the 3x3 filter. "
    "10 accumulators, 13 registers, exactly the A2 budget. Peak is 32 FLOPs per cycle because floats give 8 lanes instead of 4. "
    "10 accumulators is the minimum (5 x 2); the 2x6 tile uses 12, two chains of slack.")
rows = [("", "Assignment 2", "Stride-1 microkernel"), ("m (broadcast)", "5 rows of A", "5 output pixels"),
        ("n (vector loads)", "8 = 2 × 4 doubles", "16 ch = 2 × 8 floats"), ("k (reduction)", "32 to 512", "C × 9 (144 to 2304)"),
        ("Accumulators", "10", "10 (2×5) or 12 (2×6)"), ("Registers", "13 of 16", "13 or 15 of 16"), ("Peak per core", "16 FLOPs/cycle", "32 FLOPs/cycle")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.5), Inches(1.2), Inches(5.0), Inches(2.6)).table
for c, wd in enumerate((1.45, 1.6, 1.95)): tbl.columns[c].width = Inches(wd)
for r, row in enumerate(rows):
    for c, v in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = v
        cell.fill.solid(); cell.fill.fore_color.rgb = SOFT if r == 0 else WHITE
        cell.margin_left = cell.margin_right = Inches(0.06)
        f = cell.text_frame.paragraphs[0].runs[0].font if v else None
        if f: f.size, f.name, f.bold, f.color.rgb = Pt(11), FONT, r == 0 or c == 0 and False, INK
text(s, 0.5, 3.95, 5.0, 0.6, "Same Broadwell core: 2 FMA units × 8 lanes × 2 FLOPs = 32 FLOPs/cycle. Latency 5 × throughput 2 = 10 chains in flight.", size=12, color=MUTED)
for i, (h, b) in enumerate([("C", "input channels of the layer (16 to 256)"),
                            ("k = C × 9", "multiply-adds per output: C channels × the 9 taps of the 3×3 filter (A2's k)"),
                            ("accumulator", "one YMM register = 8 outputs (8 channels of one pixel); one dependency chain"),
                            ("tile 2×5 (10)", "outputs per call: 2 blocks of 8 ch × 5 pixels = 80 outputs in 10 accumulators")]):
    y = 1.2 + i * 0.84
    card(s, 5.75, y, 3.75, 0.76, PINK if i == 3 else SOFT, f"term-{i}")
    text(s, 5.88, y + 0.07, 3.5, 0.25, h, size=12, bold=True, color=RED if i == 3 else INK, font="Courier New")
    text(s, 5.88, y + 0.33, 3.5, 0.42, b, size=10, color=MUTED)

# ------------------------------------------------------------------ 3. tile as tensors
s = content_slide("One kernel call = 80 outputs",
    "One call computes the teal strip: one output row, 5 pixels wide, 16 output channels deep, 80 numbers. "
    "Each cell is 8 channels of one pixel, exactly one YMM register, so the strip is 10 accumulators. "
    "It reads a 3 by 7 window of the input across all C channels and 16 filters of 3 by 3 by C. "
    "Each step multiplies 2 weight vectors by 5 broadcast inputs, 10 FMAs, repeated k = C x 9 times, then bias, SiLU, store.")
s.shapes.add_picture(fig("docs/milestone/om/figs/tile_tensors.png"), Inches(1.2), Inches(1.05), width=Inches(7.6))
text(s, 0.5, 3.3, 9.0, 0.3, "Per step: 2 weight loads + 5 broadcasts feed 10 FMAs; repeated k = C × 9 times, then bias, SiLU, one store.",
     size=11.5)
strip = [("SIMD", "one vfmadd231ps = 8 multiply-adds (8 output channels): 32 FLOPs/cycle vs 4 scalar"),
         ("Intrinsics", "C calls that are one instruction each, e.g. _mm256_fmadd_ps = vfmadd231ps"),
         ("Reuse", "each weight vector used 5×, each broadcast 2×, each accumulator stored once"),
         ("Layout", "8 channels of a pixel are contiguous: one aligned load fills one register")]
for i, (h, b) in enumerate(strip):
    x = 0.5 + i * 2.28
    card(s, x, 3.68, 2.15, 0.92, PINK if i == 0 else SOFT, f"strip-{i}")
    text(s, x + 0.1, 3.74, 1.95, 0.25, h, size=11.5, bold=True, color=RED if i == 0 else INK)
    text(s, x + 0.1, 3.99, 1.95, 0.6, b, size=9.5, color=MUTED)

# ------------------------------------------------------------------ 3c. design questions (a)-(c)
s = content_slide("(a)–(c): size, registers, allocation",
    "Answering the handout questions directly. (a) One call computes 80 outputs: 5 pixels of one output row times 16 output channels. "
    "(b) It needs 13 YMM registers out of the 16 that AVX2 gives on Broadwell: 10 accumulators, 2 weight vectors, 1 broadcast input. "
    "10 accumulators is the minimum to hide FMA latency: 5 cycles latency times 2 FMA units. "
    "(c) Outputs: the 10 accumulators stay in registers for the whole C x 9 reduction. Inputs: the 2 weight registers hold 8 output channels each for one input channel and tap; "
    "the broadcast register holds one input value copied to all 8 lanes. Temporaries: none inside the loop; 3 registers stay free. "
    "We write intrinsics, so GCC assigns the physical registers; we checked GCC's output for v2 and it keeps all 13 in registers with no spills.")
# register map: 16 YMM boxes
regs = ([("c[0][%d]" % j, TEALSOFT, TEAL) for j in range(5)] + [("c[1][%d]" % j, TEALSOFT, TEAL) for j in range(5)]
        + [("w0", AMBER_SOFT, AMBER), ("w1", AMBER_SOFT, AMBER), ("x", RGBColor(0xE4, 0xE0, 0xF0), RGBColor(0x5B, 0x4B, 0x8A))]
        + [("free", WHITE, GREY)] * 3)
text(s, 0.5, 1.05, 4.7, 0.3, "16 YMM registers, 8 floats each (row 1: block kb, row 2: kb+1)", size=11, bold=True)
for i, (lab, fc, ec) in enumerate(regs):
    row, col = (i // 5, i % 5) if i < 10 else (2, i - 10)
    bw = 0.86 if i < 10 else 0.71
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5 + col * (bw + 0.06)), Inches(1.45 + row * 0.62), Inches(bw), Inches(0.52))
    b.fill.solid(); b.fill.fore_color.rgb = fc; b.line.color.rgb = ec; b.line.width = Pt(1.25); b.shadow.inherit = False
    if lab == "free": b.line.dash_style = 4
    tf = b.text_frame; tf.margin_left = tf.margin_right = 0
    r = tf.paragraphs[0].add_run(); r.text = lab
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r.font.size, r.font.name, r.font.color.rgb, r.font.bold = Pt(11), "Courier New", ec, lab != "free"
text(s, 0.5, 3.45, 4.6, 0.9, [[("c[b][j]", {"font": "Courier New", "bold": True, "color": TEAL}), (" output block b (8 ch), pixel j: 10 accumulators", {})],
                            [("w0, w1", {"font": "Courier New", "bold": True, "color": AMBER}), (" 8 weights for blocks kb, kb+1 (one load each)", {})],
                            [("x", {"font": "Courier New", "bold": True, "color": RGBColor(0x5B, 0x4B, 0x8A)}), (" one input value broadcast to 8 lanes", {})]],
     size=10.5, color=MUTED)
qa = [("(a) Size", "80 outputs per call: 5 pixels × 16 output channels (2 blocks of 8)."),
      ("(b) Registers", "Need 13 of the 16 YMM registers (AVX2, no AVX-512). At least 10 accumulators: FMA latency 5 × 2 FMA units."),
      ("(c) Allocation", "Outputs: 10 accumulators, live for the whole C × 9 loop. Inputs: 2 weight + 1 broadcast. Temporaries: none in the loop; 3 free. Pointers and counters live in integer registers.")]
for i, (h, b) in enumerate(qa):
    y = 1.05 + i * 1.2
    card(s, 5.25, y, 4.25, 1.1, PINK if i == 1 else SOFT, f"qa-{i}")
    text(s, 5.38, y + 0.08, 4.0, 0.28, h, size=12.5, bold=True, color=RED if i == 1 else INK)
    text(s, 5.38, y + 0.38, 4.0, 0.7, b, size=10.5, color=MUTED)

# ------------------------------------------------------------------ 3d. design questions (d)-(e)
s = content_slide("(d)–(e): register reuse and the code",
    "(d) Reuse. Each weight register is loaded once and used by 5 FMAs, one per pixel. Each broadcast input is used by 2 FMAs, one per output block. "
    "The single x register is reused 5 times per step, once per pixel, because each broadcast is consumed right away. "
    "Each accumulator is reused C x 9 times and written to memory once. So one step is 7 loads for 10 FMAs, which the 2 load ports handle with room to spare. "
    "(e) The code: initialize the accumulators with the bias, so the add is free. Then loop over input-channel blocks, the 3 by 3 taps, and the 8 input channels of the block; "
    "load 2 weight vectors, then for each of the 5 pixels broadcast one input and issue 2 FMAs. After the loop, apply SiLU in registers and store. "
    "The zero halo means no edge branches. GCC fully unrolls the inner 8 channels and 5 pixels, so the loop body is 80 FMAs, 40 broadcasts and 16 loads.")
reuse = [("Weight w0, w1", "1 load → 5 FMAs (one per pixel)"),
         ("Input x", "1 broadcast → 2 FMAs (one per block); the one x register is reused for all 5 pixels"),
         ("Accumulator c", "C × 9 FMAs → 1 store (bias as start value)"),
         ("Per step", "7 loads : 10 FMAs, under the 2 loads/cycle limit")]
text(s, 0.5, 1.05, 3.6, 0.3, "(d) Reuse", size=13, bold=True, color=RED)
for i, (h, b) in enumerate(reuse):
    y = 1.4 + i * 0.8
    card(s, 0.5, y, 3.6, 0.72, SOFT, f"reuse-{i}")
    text(s, 0.62, y + 0.06, 3.4, 0.25, h, size=11.5, bold=True)
    text(s, 0.62, y + 0.31, 3.4, 0.42, b, size=9.5, color=MUTED)
text(s, 4.35, 1.05, 5.1, 0.3, "(e) Implementation, AVX2 intrinsics (v2)", size=13, bold=True, color=RED)
code = [
    "c[b][j] = _mm256_load_ps(bias + 8*(kb+b));  // b<2, j<5",
    "for cb in C/8:  for r in 0..2:  for s in 0..2:",
    "  for ci in 0..7:              // unrolled by GCC",
    "    w0 = _mm256_load_ps(W[kb  ][cb][r][s][ci]);",
    "    w1 = _mm256_load_ps(W[kb+1][cb][r][s][ci]);",
    "    for j in 0..4:             // unrolled",
    "      x = _mm256_broadcast_ss(&In[cb][oh+r][ow+j+s][ci]);",
    "      c[0][j] = _mm256_fmadd_ps(x, w0, c[0][j]);",
    "      c[1][j] = _mm256_fmadd_ps(x, w1, c[1][j]);",
    "c[b][j] = epilogue_ps(c[b][j], fuse_silu);   // SiLU",
    "_mm256_store_ps(&Out[kb+b][oh+1][ow+j+1][0], c[b][j]);",
]
card(s, 4.35, 1.4, 5.15, 2.55, RGBColor(0xF7, 0xF8, 0xFA), "code")
text(s, 4.47, 1.5, 4.95, 2.4, code, size=8.5, font="Courier New", color=INK)
text(s, 4.35, 4.03, 5.15, 0.55, "Weights packed [K/8][C/8][3][3][8 ic][8 oc]; activations NCHW8c with a 1-pixel zero halo, so no edge branches. "
     "Outer loops: output-channel pairs kb, rows oh, pixel groups ow of 5.", size=9.5, color=MUTED)

# ------------------------------------------------------------------ 3b. assumptions
s = content_slide("How the numbers are measured",
    "These are the assumptions behind every number in the talk. The important one: cycles are wall time times 3.4 GHz, the CPU's maximum turbo, "
    "exactly the Assignment 2 convention. The core can never run faster than that, so the percentages can only understate.")
assume = [("Machine and code", "ece028, Xeon E5-2640 v4, 1 core (taskset); C++ with AVX2 intrinsics only, GCC 8.5 -O3"),
          ("Peak", "2 FMA units × 8 FP32 lanes × 2 FLOPs = 32 FLOPs/cycle per core (AVX2, no AVX-512)"),
          ("Cycles", "wall time × 3.4 GHz (max turbo), as in Assignment 2: can only understate % of peak"),
          ("FLOPs", "2 × outputs × C × 9: only the convolution's multiply-adds; bias and SiLU count as zero"),
          ("Microkernel alone", "one tile of data reused (stays in L1/L2), 2 s warm-up, best of 3 runs, SiLU off"),
          ("Full layers", "all 32 real YOLOv8n layers, batch 1, 640 × 640 image, FP32, SiLU on"),
          ("Correct", "every result checked: max relative error ≤ 1e-4 vs a scalar reference / OpenCV's real outputs"),
          ("Baseline", "OpenCV DNN on the same ONNX model, 1 thread, its own per-layer timer")]
for i, (h, b) in enumerate(assume):
    col, row = i % 2, i // 2
    x, y = 0.5 + col * 4.55, 1.1 + row * 0.88
    card(s, x, y, 4.4, 0.78, SOFT, f"assume-{i}")
    text(s, x + 0.12, y + 0.07, 4.1, 0.25, h, size=11.5, bold=True, color=RED if h == "Cycles" else INK)
    text(s, x + 0.12, y + 0.33, 4.15, 0.45, b, size=9.5, color=MUTED)

# ------------------------------------------------------------------ 4. four steps
s = content_slide("Four steps: ops → accumulators → tile → benchmark",
    "Step 1: every output is an independent dot product; the only chain is the reduction. "
    "Step 2: 5-cycle latency times 2 FMAs per cycle means 10 independent accumulators, 80 outputs in flight. The chart: 1 accumulator is 10%, each adds about 10 points. "
    "Step 3: 2x5 gives 10 accumulators in 13 registers, 7 loads per 10 FMAs, and 5 divides 160, 80, 40 and 20. "
    "Step 4: benchmark the kernel alone exactly like Assignment 2.")
steps = [("Independent ops", "Each output is a dot product of length C × 9: 102,400 in a 40×40, 64-ch layer."),
         ("Outputs in flight", "5 × 2 FMA/cycle = 10 accumulators = 80 outputs; 1 accumulator = 10% of peak."),
         ("Kernel size", "2×5 tile: 10 acc + 2 weights + 1 broadcast = 13 regs; 7 loads per 10 FMAs."),
         ("Benchmark like A2", "Kernel alone, k = 72 to 4608, 3 runs, 2 s warm-up, cycles at 3.4 GHz, checked.")]
for i, (h, b) in enumerate(steps):
    y = 1.15 + i * 0.85
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.5), Inches(y + 0.02), Inches(0.36), Inches(0.36))
    o.fill.solid(); o.fill.fore_color.rgb = RED; o.line.fill.background(); o.shadow.inherit = False
    text(s, 0.5, y + 0.02, 0.36, 0.36, str(i + 1), size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.0, y - 0.02, 3.6, 0.28, h, size=13, bold=True)
    text(s, 1.0, y + 0.27, 3.6, 0.52, b, size=10.5, color=MUTED)
cd = CategoryChartData()
cd.categories = ["1×1 (1)", "2×2 (4)", "1×5 (5)", "2×4 (8)", "1×10 (10)", "2×5 (10)"]
cd.add_series("% of peak", (9.9, 39.3, 49.6, 72.4, 85.8, 85.3))  # results/ukr_C64_nosilu_34.csv
ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4.8), Inches(1.05), Inches(4.7), Inches(3.55), cd).chart
style_chart(ch, '0"%"', TEAL)
ch.value_axis.maximum_scale, ch.value_axis.minimum_scale, ch.value_axis.major_unit = 100, 0, 25
ch.value_axis.tick_labels.number_format, ch.value_axis.tick_labels.number_format_is_linked = '0"%"', False
chart_title(ch, "Microkernel alone vs accumulators (C = 64, SiLU off)")

# ------------------------------------------------------------------ 5. design space
s = content_slide("Same math, 50% to 86%: the tile and GCC's registers",
    "Every kernel I built, all intrinsics, same arithmetic, same layout. Filled dots: the microkernel alone; hollow: all 32 layers. "
    "2x5 is the best in the layers: 85% alone, 77% in all 32 layers. "
    "Bigger tiles lose because of how GCC allocates registers: at 15 live registers (2x6) it adds 11 register copies per tap, found by reading GCC's generated assembly; "
    "at 16 (4x3) it spills accumulators to memory and drops to 49%. "
    "Unrolling and software pipelining made it worse: full unrolling overflowed the uop cache and partial unrolling made GCC spill.")
s.shapes.add_picture(fig("results/ukr_design_space_intrinsics.png"), Inches(0.5), Inches(1.1), height=Inches(3.5))
for i, (h, b) in enumerate([("Best: 2×5", "13 registers, 85% alone, 77% in all 32 layers"),
                            ("GCC's registers", "clean at 13, 11 copies per tap at 15 (2×6), spills at 16 (4×3: 49%)"),
                            ("Unroll, pipeline", "lost: uop cache overflow, spills")]):
    y = 1.1 + i * 1.18
    card(s, 7.05, y, 2.45, 1.05, PINK if i == 0 else SOFT, f"ds-card-{i}")
    text(s, 7.17, y + 0.08, 2.2, 0.28, h, size=12, bold=True, color=RED if i == 0 else INK)
    text(s, 7.17, y + 0.38, 2.2, 0.65, b, size=10, color=MUTED)

# ------------------------------------------------------------------ 6. sweep
s = content_slide("Benchmarked like Assignment 2: FLOPs/cycle vs k",
    "This is the Assignment 2 figure for my kernel: FLOPs per cycle against the reduction length k. "
    "Cycles are time times 3.4 GHz, the max turbo, exactly as in Assignment 2, so this can only understate. "
    "2x5 holds 27 to 28.5 FLOPs per cycle, 83 to 89 percent of peak; 2x6 sits about 3 points lower, 80 to 83, because of GCC's register copies. "
    "The dip near k = 500 is the 2x5 weights leaving L1, like k = 315 in A2. "
    "Cross-checked: the hardware FP-instruction counter matches our FLOP count exactly.")
s.shapes.add_picture(fig("results/ukr_sweep_nosilu_34_intrinsics.png"), Inches(0.5), Inches(1.05), height=Inches(3.2))
text(s, 0.5, 4.3, 6.0, 0.3, "SiLU off, cycles = time × 3.4 GHz, as in A2. 3 runs, every point checked (576 pass).", size=10, color=MUTED)
stat(s, 6.75, 1.05, 2.75, "83–89%", "2×5 (v2's tile): 27–28.5 FLOPs/cycle, ~93 GFLOP/s on one core", RED)
stat(s, 6.75, 2.2, 2.75, "80–83%", "2×6: GCC adds register copies at 15 live registers", GREY)
stat(s, 6.75, 3.35, 2.75, "k ≈ 500", "L1 dip: 2×5 weights pass 32 KB, like k ≈ 315 in A2", AMBER)

# ------------------------------------------------------------------ 7. alone vs in a layer
s = content_slide("Tied alone, 10 points apart in the layers",
    "Two tiles with the same 10 accumulators: 1x10, one block of 8 channels times 10 pixels, and 2x5. Alone they tie: 85.8 and 85.3 percent. "
    "Inside real layers 2x5 wins by 10 points: 76.8 against 67.0. "
    "Alone, everything sits in L1, so 1x10's 11 loads per 10 FMAs still fit through the 2 load ports. "
    "In a layer the inputs and weights stream from L2, and because each 1x10 call covers only 8 output channels, the layer reads every input row twice as often as with 2x5. "
    "So loads become the limit. 2x5 needs only 7 loads per 10 FMAs, so it has room to spare. Lesson: pick the tile in the real layer, not only alone.")
cd = CategoryChartData()
cd.categories = ["1×10 (10 acc)", "2×5 (10 acc)"]
cd.add_series("alone", (85.8, 85.3))
cd.add_series("in layers", (67.0, 76.8))
gc = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.1), Inches(4.4), Inches(3.5), cd)
ch = gc.chart
style_chart(ch, '0.0"%"', TEAL)
ch.plots[0].series[1].format.fill.solid(); ch.plots[0].series[1].format.fill.fore_color.rgb = GREY
ch.has_legend = True
from pptx.enum.chart import XL_LEGEND_POSITION
ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
ch.legend.font.size, ch.legend.font.name = Pt(10), FONT
ch.value_axis.maximum_scale, ch.value_axis.minimum_scale, ch.value_axis.major_unit = 100, 0, 25
ch.value_axis.tick_labels.number_format, ch.value_axis.tick_labels.number_format_is_linked = '0"%"', False
chart_title(ch, "% of peak: alone (C = 64) vs in layers")
rows2 = [("", "1×10", "2×5"), ("Loads per 10 FMAs", "11", "7"), ("Output channels per call", "8", "16"),
         ("Times each input row is read", "2×", "1×")]
tb = s.shapes.add_table(len(rows2), 3, Inches(5.2), Inches(1.15), Inches(4.3), Inches(1.5)).table
for c, wd in enumerate((2.3, 1.0, 1.0)): tb.columns[c].width = Inches(wd)
for r, row in enumerate(rows2):
    for c, v in enumerate(row):
        cell = tb.cell(r, c); cell.text = v
        cell.fill.solid(); cell.fill.fore_color.rgb = SOFT if r == 0 else WHITE
        if v:
            f = cell.text_frame.paragraphs[0].runs[0].font
            f.size, f.name, f.color.rgb = Pt(11), FONT, INK
            f.bold = r == 0
text(s, 5.2, 2.85, 4.3, 0.9, [[("Alone: ", {"bold": True, "color": INK}), ("all data in L1, so 1×10's extra loads still fit.", {})],
                             [("In a layer: ", {"bold": True, "color": INK}), ("data streams from L2 and 1×10 reads inputs twice as often, so loads become the limit.", {})]],
     size=11.5, color=MUTED)
text(s, 5.2, 3.95, 4.3, 0.6, "Pick the tile in the real layer, not only alone.", size=14, bold=True, italic=True, color=RED)
text(s, 5.2, 4.5, 4.3, 0.2, "1×10 in layers: 4-layer tile sweep (C = 16–128).", size=9, color=MUTED)

# ------------------------------------------------------------------ 8. against OpenCV
s = content_slide("1.24× faster than OpenCV on all 32 layers",
    "Total time for all 32 stride-1 layers, one core, same model and image. OpenCV DNN takes 86.4 milliseconds. "
    "Our kernel, v2, the 2x5 tile written with intrinsics, takes 69.5: 1.24 times faster, at 77 percent of peak against OpenCV's 62. "
    "Next: cache blocking for the layers with 128 or more input channels, and threading across 20 cores.")
cd = CategoryChartData()
cd.categories = ["v2 (2×5, intrinsics)", "OpenCV DNN"]
cd.add_series("ms", (69.5, 86.4))
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.1), Inches(5.6), Inches(3.4), cd).chart
style_chart(ch, '0.0" ms"', GREY)
ch.value_axis.maximum_scale, ch.value_axis.minimum_scale, ch.value_axis.major_unit = 100, 0, 25
chart_title(ch, "All 32 layers, 1 core, SiLU on: ms (lower is better)")
stat(s, 6.5, 1.1, 3.0, "1.24×", "v2 vs OpenCV DNN, all 32 layers, 1 core", RED)
stat(s, 6.5, 2.25, 3.0, "77% vs 62%", "of peak: v2 vs OpenCV DNN", GREY)
text(s, 6.5, 3.45, 3.0, 1.1, "Next: cache blocking for C ≥ 128, then threads on 20 cores (v3: 13.5× on 20 cores).", size=11.5, italic=True, color=MUTED)

# ------------------------------------------------------------------ backup
s = content_slide("Backup: the same sweep with SiLU on",
    "With the fused SiLU on, the curves sit about 4 to 5 points lower: SiLU costs a fixed 160 to 290 cycles per call, a big share at small k and almost nothing from C = 128 up.")
s.shapes.add_picture(fig("results/ukr_sweep_silu_34_intrinsics.png"), Inches(0.5), Inches(1.05), height=Inches(3.2))
text(s, 6.75, 1.1, 2.75, 3.3, [[("SiLU on: ", {"bold": True, "color": INK}), ("2×5 81.4% mean (C = 32–256), 2×6 about 3 points lower.", {})],
                              [("Cost: ", {"bold": True, "color": INK}), ("a fixed 160–290 cycles per call, 20% of the call at C = 16, about 4.5% at C = 64, nothing measurable from C = 128.", {})]],
     size=11.5, color=MUTED)

prs.save(OUT)
print("wrote", OUT)
